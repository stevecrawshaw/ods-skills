"""Start a WECA deck from the official PowerPoint template.

The template ships 27 layouts and 13 example slides. Opening it with
python-pptx and calling add_slide() appends after those examples, so strip them
first. Every layout already carries the correct logo (white on colour, black on
white), so never add a logo picture yourself.

Usage:
    from weca_pptx import open_deck, add
    prs = open_deck()
    s = add(prs, "West green title slide A", title="Local Transport Plan",
            placeholders={10: "September 2026"})
    prs.save("deck.pptx")

Run this file to print every layout with its placeholder idx values.
"""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.presentation import Presentation as PptxPresentation
from pptx.slide import Slide

TEMPLATE = Path(__file__).resolve().parent.parent / "reference" / "slide_template_weca.pptx"


def open_deck(path: Path = TEMPLATE) -> PptxPresentation:
    """Open the template with its 13 example slides removed."""
    prs = Presentation(str(path))
    sld_ids = prs.slides._sldIdLst
    for sld_id in list(sld_ids):
        prs.part.drop_rel(sld_id.rId)
        sld_ids.remove(sld_id)
    return prs


def add(
    prs: PptxPresentation,
    layout: str,
    title: str | None = None,
    placeholders: dict[int, str] | None = None,
) -> Slide:
    """Add a slide by layout name (names have stray spaces; matched stripped)."""
    layouts = {lay.name.strip(): lay for lay in prs.slide_layouts}
    if layout not in layouts:
        raise KeyError(f"{layout!r} not in template; choose from {sorted(layouts)}")
    slide = prs.slides.add_slide(layouts[layout])
    if title is not None and slide.shapes.title is not None:
        slide.shapes.title.text = title
    for idx, text in (placeholders or {}).items():
        slide.placeholders[idx].text = text
    return slide


if __name__ == "__main__":
    # idx numbers do NOT follow reading order (e.g. on "Three-column text" idx 11
    # is the subtitle and idx 10 the first column), so list them top-to-bottom,
    # left-to-right with their position in inches.
    prs = open_deck()
    for i, lay in enumerate(prs.slide_layouts):
        print(f"{i:2d} {lay.name.strip()}")
        for p in sorted(lay.placeholders, key=lambda p: (round(p.top / 914400, 1), p.left)):
            f = p.placeholder_format
            print(
                f"     idx {f.idx:2d} {f.type.name.lower():12s} "
                f"at ({p.left / 914400:.1f}, {p.top / 914400:.1f}) in, "
                f"{p.width / 914400:.1f} wide"
            )
